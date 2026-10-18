"""Build a portable three-level atlas from saved tracks, never from raw contacts.

Original PNGs and numerical runs are read-only. The default interactive and new
static views use a pinned RefSeq gene snapshot; original author gene coordinates
remain explicitly distinguished in the retained PNG view.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
import csv
from datetime import datetime, timezone
import gzip
import hashlib
import html
import json
from pathlib import Path
import re
import shutil
from urllib.parse import unquote

import numpy as np

ROOT = Path(__file__).resolve().parents[3]
CONFIG = ROOT / "configs/task3/atlas/atlas_v1.json"
TYPES = ("OPCID", "CHIN", "CHID")
COLORS = ("#2463a7", "#cf8427", "#a53e6b")


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def rows(path):
    with Path(path).open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


def write_csv(path, values):
    with Path(path).open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(values[0]))
        writer.writeheader()
        writer.writerows(values)


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n")


def record(path, root=ROOT):
    return {"path": str(Path(path).relative_to(root)), "bytes": Path(path).stat().st_size, "sha256": sha(path)}


def read_reference(path, chrom, length):
    """GFF is 1-based closed. Convert each gene to 0-based half-open exactly once."""
    raw = Path(path).read_bytes()
    text = gzip.decompress(raw).decode() if raw.startswith(b"\x1f\x8b") else raw.decode()
    region = f"##sequence-region {chrom} 1 {length}"
    if region not in text.splitlines():
        raise ValueError("Reference sequence identity/length mismatch")
    result = []
    for line in text.splitlines():
        if not line or line.startswith("#"):
            continue
        a = line.split("\t")
        if len(a) != 9 or a[2] not in ("gene", "pseudogene"):
            continue
        if a[0] != chrom:
            raise ValueError("Unexpected reference chromosome")
        attr = {k: unquote(v) for k, v in (x.split("=", 1) for x in a[8].split(";") if "=" in x)}
        start, end = int(a[3]) - 1, int(a[4])
        if not 0 <= start < end <= length or a[6] not in ("+", "-"):
            raise ValueError("Invalid reference gene interval or strand")
        result.append({"reference_id": attr["ID"], "locus_tag": attr.get("locus_tag", ""),
                       "gene": attr.get("gene", attr.get("Name", "")), "chrom": chrom,
                       "start": start, "end": end, "strand": a[6],
                       "feature_type": a[2], "biotype": attr.get("gene_biotype", ""),
                       "aliases": attr.get("gene_synonym", ""), "coordinate_system": "0-based-half-open"})
    if not result or len({r["reference_id"] for r in result}) != len(result):
        raise ValueError("Empty or nonunique reference gene IDs")
    return sorted(result, key=lambda r: (r["start"], r["end"], r["reference_id"]))


def audit_gene_rows(source, reference):
    lookup = defaultdict(list)
    for r in reference:
        for name in set([r["gene"], r["locus_tag"]] + r["aliases"].split(",")) - {""}:
            lookup[name].append(r)
    result = []
    for i, s in enumerate(source, 2):
        candidates = lookup[s["gene"]]
        start, end = int(s["start"]), int(s["end"])
        half = [r for r in candidates if (r["start"], r["end"], r["strand"]) == (start, end, s["strand"])]
        closed = [r for r in candidates if (r["start"] + 1, r["end"], r["strand"]) == (start, end, s["strand"])]
        status = ("exact_0_halfopen" if half else "exact_1_closed" if closed else
                  "name_match_coordinates_differ" if candidates else "name_not_found")
        matched = half or closed
        result.append({"source_row": i, "gene": s["gene"], "source_start": start, "source_end": end,
                       "source_strand": s["strand"], "status": status,
                       "matched_reference_ids": ";".join(r["reference_id"] for r in matched),
                       "name_candidate_reference_ids": ";".join(r["reference_id"] for r in candidates),
                       "display_action": "use_separate_pinned_RefSeq_layer; preserve_source_row"})
    return result


def check_pages(pages, directory, length, width):
    """Validate actual images and exact tiling, not just a 465-row index."""
    from PIL import Image
    if len(pages) != (length + width - 1) // width:
        raise ValueError("Missing track pages")
    seen = set()
    for i, r in enumerate(pages):
        start, end = i * width, min((i + 1) * width, length)
        expected = f"tracks/{start:07d}_{end:07d}.png"
        if (int(r["start"]), int(r["end"]), r["file"]) != (start, end, expected):
            raise ValueError("Noncontiguous or incorrect page coordinate/path")
        image = Path(directory) / expected
        if image.is_symlink() or not image.is_file():
            raise ValueError(f"Missing/linked track image: {expected}")
        with Image.open(image) as im:
            if im.format != "PNG" or min(im.size) < 500:
                raise ValueError("Invalid track image")
            im.verify()
        seen.add(Path(expected).name)
    if {p.name for p in (Path(directory) / "tracks").glob("*.png")} != seen:
        raise ValueError("Unindexed track images")


def validate_run(directory, config):
    directory = Path(directory)
    manifest = json.loads((directory / "manifest.json").read_text())
    for r in manifest["output_files"]:
        path = Path(r["path"])
        if not path.resolve().is_relative_to(directory.resolve()) or path.is_symlink():
            raise ValueError("Unsafe run-manifest output")
        if path.stat().st_size != r["size_bytes"] or sha(path) != r["sha256"]:
            raise ValueError(f"Run output signature differs: {path}")
    pages = rows(directory / "task3/pages.csv")
    check_pages(pages, directory / "task3", config["genome_length_bp"], config["page_width_bp"])
    with np.load(directory / "task3/signals.npz", allow_pickle=False) as z:
        tracks = {k: z[k].copy() for k in z.files}
    n = (config["genome_length_bp"] + 99) // 100
    if tracks["signals"].shape != (6, n) or not np.array_equal(tracks["positions"], np.arange(n) * 100):
        raise ValueError("Signal coordinates/coverage differ")
    valid = tracks["common_valid"]
    if valid.dtype != bool or not np.isfinite(tracks["signals"][:, valid]).all() or not np.isnan(tracks["signals"][:, ~valid]).all():
        raise ValueError("Missing values or common-valid mask differ")
    np.testing.assert_allclose(tracks["condition_means"], tracks["signals"].reshape(3, 2, n).mean(axis=1), equal_nan=True)
    return pages, tracks, len(manifest["output_files"])


def audit_author_sources(config, source_genes, structures):
    import openpyxl
    wb = openpyxl.load_workbook(ROOT / config["author_workbook"], read_only=True, data_only=True)
    try:
        original_genes = [{"gene": r[0], "start": str(int(r[2])), "end": str(int(r[3])), "strand": r[4]}
                          for r in wb["Supplementary Table 1"].iter_rows(min_row=2, values_only=True) if r[0]]
        if original_genes != [{k: r[k] for k in original_genes[0]} for r in source_genes]:
            raise ValueError("Source gene table differs from workbook rows")
        result = []
        by_sheet = {}
        for r in structures:
            sheet, line = r["source_sheet"], int(r["source_row"])
            if sheet not in by_sheet:
                by_sheet[sheet] = list(wb[sheet].iter_rows(values_only=True))
            values = by_sheet[sheet][line - 1]
            end_index = 4 if r["type"] == "CHIN" else 3
            if (values[0], int(values[2]), int(values[end_index])) != (r["structure_id"], int(r["start"]), int(r["end"])):
                raise ValueError("Author structure differs from original source")
            if not 0 <= int(r["start"]) < int(r["end"]) <= config["genome_length_bp"]:
                raise ValueError("Author interval outside reference")
            result.append({"structure_id": r["structure_id"], "type": r["type"], "source_sheet": sheet,
                           "source_row": line, "start": r["start"], "end": r["end"],
                           "source_values_verified": True, "coordinate_convention": "provisional; source values retained",
                           "display_action": "author_track_separate_from_RefSeq; do_not_relabel"})
        return result
    finally:
        wb.close()


def plot_views(config, tracks, structures, reference, output):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import Rectangle
    from matplotlib.ticker import FuncFormatter
    pos = tracks["positions"]
    signals = np.where(tracks["common_valid"][None, :], tracks["signals"], np.nan)
    means = signals.reshape(3, 2, -1).mean(axis=1)
    for detail, name in [(False, "whole_genome.png"), (True, "representative_region.png")]:
        lo, hi = config["representative_region_bp"] if detail else (0, config["genome_length_bp"])
        selected = (pos >= lo) & (pos < hi)
        x = pos[selected]
        ys = signals[:, selected]
        av = means[:, selected]
        if not detail:
            x = np.array([x[i:i + 100].mean() for i in range(0, len(x), 100)])
            ys = np.array([[np.nanmean(y[i:i + 100]) if np.isfinite(y[i:i + 100]).any() else np.nan for i in range(0, len(y), 100)] for y in ys])
            av = ys.reshape(3, 2, -1).mean(axis=1)
        divisor = 1_000 if detail else 1_000_000
        fig, axes = plt.subplots(6, 1, figsize=(13.5, 10.5), sharex=True,
                                 gridspec_kw={"height_ratios": [1.5, 1.5, 1.5, 1.25, 1.2, .22]})
        fig.subplots_adjust(left=.115, right=.975, top=.86, bottom=.13, hspace=.25)
        ymax = float(np.nanmax(np.sqrt(ys))) * 1.08
        for j, ax in enumerate(axes[:3]):
            for k in (2*j, 2*j + 1):
                ax.plot(x/divisor, np.sqrt(ys[k]), color=COLORS[j], alpha=.35, lw=.6)
            ax.plot(x/divisor, np.sqrt(av[j]), color=COLORS[j], lw=1.05)
            ax.set(ylim=(0, ymax), ylabel=("WT", r"$\Delta stpA$", r"$\Delta hns\Delta stpA$")[j])
            if not detail:
                ax.axvspan(1.05, 1.10, color="#334155", alpha=.10)
            ax.grid(axis="y", color="#e3e8ee", lw=.6)
        label_counts = Counter()
        for r in structures:
            a, z = int(r["start"]), int(r["end"])
            if a < hi and z > lo:
                j = TYPES.index(r["type"])
                axes[3].add_patch(Rectangle((max(a, lo)/divisor, j-.24), (min(z, hi)-max(a, lo))/divisor, .48, facecolor=COLORS[j], alpha=.65))
                if detail:
                    label_counts[r["type"]] += 1
                    axes[3].text((max(a, lo)+min(z, hi))/2/divisor, j+.27+.22*(label_counts[r["type"]] % 2), r["structure_id"], fontsize=6, ha="center", clip_on=True)
        axes[3].set(ylim=(2.65, -.55), yticks=(0, 1, 2), yticklabels=TYPES, ylabel="Author\nstructures")
        genes = [r for r in reference if r["start"] < hi and r["end"] > lo]
        occupied = {"+": -1, "-": -1}
        for r in genes:
            j = 0 if r["strand"] == "+" else 1
            a, z = max(r["start"], lo), min(r["end"], hi)
            axes[4].add_patch(Rectangle((a/divisor, j-.20), (z-a)/divisor, .40, facecolor="#647f91", alpha=.7))
            if detail:
                direction = 1 if j == 0 else -1
                axes[4].annotate("", xy=((z if direction == 1 else a)/divisor, j), xytext=((a if direction == 1 else z)/divisor, j), arrowprops={"arrowstyle": "->", "color": "#375368", "lw": .5})
                if a >= occupied[r["strand"]]:
                    axes[4].text((a+z)/2/divisor, j+.32, r["gene"], fontsize=6, ha="center", clip_on=True)
                    occupied[r["strand"]] = a + 1050
        axes[4].set(ylim=(1.65, -.5), yticks=(0, 1), yticklabels=("+", "−"), ylabel="RefSeq\ngenes")
        mask = tracks["common_valid"][selected].astype(float)[None, :]
        from matplotlib.colors import ListedColormap
        axes[5].imshow(mask, aspect="auto", interpolation="nearest", extent=(lo/divisor, hi/divisor, 0, 1), cmap=ListedColormap(["#9ca3af", "white"]), vmin=0, vmax=1)
        axes[5].set(yticks=[], ylabel="QC", xlabel=f"Genome coordinate ({'kb' if detail else 'Mb'}); 0-based half-open")
        for ax in axes:
            ax.set_xlim(lo/divisor, hi/divisor)
            ax.spines[["top", "right"]].set_visible(False)
        title = "Level 3 | Predefined 1.05–1.10 Mb region" if detail else "Level 1 | Complete 4.64 Mb genome overview"
        fig.suptitle(title, x=.115, y=.975, ha="left", va="top", fontsize=18, fontweight="bold", color="#18314c")
        fig.text(.115, .923, "NC_000913.3 / 4,641,652 bp  ·  3 conditions / 6 biological samples\nThin = repeats; thick = condition mean. Shared y-scale; sqrt(relative contacts per valid pair).", va="top", fontsize=10, color="#475569")
        fig.text(.115, .025, f"{'100 bp detail' if detail else '10 kb display aggregation; full 100 bp values retained'} | RefSeq snapshot GCF_000005845.2: 4,651 gene/pseudogene features.\nAuthor structures retain source coordinates (provisional convention). Gray QC marks missing common-valid bins.\n{'Region predefined independently of classification performance and condition-specific effect size.' if detail else 'Shading locates the fixed representative region; compressed annotations are context, not new structure calls.'}", fontsize=8.5, color="#475569")
        fig.savefig(output/name, dpi=180, facecolor="white")
        plt.close(fig)


def build(config_path=CONFIG):
    config = json.loads(Path(config_path).read_text())
    out, report_dir = ROOT/config["output_dir"], ROOT/config["report_dir"]
    if out.exists() or report_dir.exists():
        raise FileExistsError("Use a new atlas version; existing evidence is never overwritten")
    for key in ("reference_gff", "assembly_report", "author_workbook"):
        if sha(ROOT/config[key]) != config[key+"_sha256"]:
            raise ValueError(f"Pinned reference signature mismatch: {key}")
    assembly = (ROOT/config["assembly_report"]).read_text()
    if "U00096.3\t=\tNC_000913.3\tPrimary Assembly\t4641652" not in assembly:
        raise ValueError("Assembly accession/sequence alias differs")
    run = ROOT/config["run_dir"]
    pages, tracks, checked = validate_run(run, config)
    genes = rows(ROOT/config["source_gene_table"])
    if genes != rows(run/"genes.csv"):
        raise ValueError("Run gene table differs from supplied source")
    structures = rows(ROOT/config["author_structures"])
    if len(structures) != 344 or len({r["structure_id"] for r in structures}) != 344:
        raise ValueError("Author structure identity differs")
    reference = read_reference(ROOT/config["reference_gff"], config["chrom"], config["genome_length_bp"])
    gene_audit = audit_gene_rows(genes, reference)
    structure_audit = audit_author_sources(config, genes, structures)
    out.mkdir(parents=True);report_dir.mkdir(parents=True)
    (out/"tracks").mkdir();(out/"pages").mkdir()
    write_csv(out/"reference_genes.csv", reference)
    write_csv(out/"supplied_gene_audit.csv", gene_audit)
    write_csv(out/"author_structure_audit.csv", structure_audit)
    write_csv(out/"author_structures.csv", [{k: r[k] for k in ("structure_id", "type", "chrom", "start", "end", "coordinate_status")} for r in structures])
    shutil.copy2(ROOT/config["reference_gff"], out/"reference_genomic.gff.gz")
    shutil.copy2(ROOT/config["assembly_report"], out/"assembly_report.txt")
    page_rows = []
    for i, page in enumerate(pages):
        shutil.copy2(run/"task3"/page["file"], out/page["file"])
        name = Path(page["file"]).stem + ".html"
        (out/"pages"/name).write_text(page_html(i, page))
        page_rows.append({"page": i+1, **page, "html": "pages/"+name})
    write_csv(out/"pages.csv", page_rows)
    plot_views(config, tracks, structures, reference, out)
    # Decimal presentation transport only; full-precision saved NPZ is unchanged.
    payload = {"length": config["genome_length_bp"], "step": 100, "pageWidth": config["page_width_bp"],
               "labels": config["labels"], "fixed": config["representative_region_bp"],
               "signals": [[None if not np.isfinite(v) else round(float(v), 12) for v in y] for y in tracks["signals"]],
               "genes": [{k: r[k] for k in ("gene", "locus_tag", "start", "end", "strand", "aliases")} for r in reference],
               "structures": [{"id": r["structure_id"], "type": r["type"], "start": int(r["start"]), "end": int(r["end"])} for r in structures],
               "pages": page_rows}
    (out/"atlas_data.js").write_text("window.MICROCTURE_ATLAS = "+json.dumps(payload, ensure_ascii=False, separators=(",", ":"), allow_nan=False)+";\n")
    asset_dir = Path(__file__).parent/"assets"
    for name in ("atlas.js", "atlas.css"):
        shutil.copy2(asset_dir/name, out/name)
    (out/"index.html").write_text(index_html(page_rows))
    valid = tracks["common_valid"]
    summary = {"schema_version": 1, "status": "three_level_atlas_verified", "created_utc": datetime.now(timezone.utc).isoformat(),
               "levels": ["whole_genome_overview", "465_page_complete_browser", "predefined_representative_region"],
               "genome_length_bp": config["genome_length_bp"], "samples": 6, "conditions": 3,
               "bins": len(valid), "common_valid_bins": int(valid.sum()), "pages": len(pages), "png_images": len(pages),
               "page_coordinate_system": "0-based-half-open", "last_page_bp": [int(pages[-1]["start"]), int(pages[-1]["end"])],
               "reference_gene_features": len(reference), "reference_feature_types": dict(Counter(r["feature_type"] for r in reference)),
               "supplied_gene_rows": len(genes), "unique_supplied_gene_names": len({r["gene"] for r in genes}),
               "supplied_gene_coordinate_comparison": dict(Counter(r["status"] for r in gene_audit)),
               "author_structures_verified_against_workbook": len(structure_audit), "author_structure_types": dict(Counter(r["type"] for r in structures)),
               "author_workbook_publisher_sha256_verified": True,
               "author_workbook_public_url": config["author_workbook_public_url"],
               "author_workbook_sha256": config["author_workbook_sha256"],
               "author_coordinate_convention": "provisional; original source values unchanged",
               "reference_layer": "pinned NCBI RefSeq GFF gene+pseudogene features, converted 1-based closed to 0-based half-open",
               "legacy_png_layer": "Original 465 PNGs retain supplied author gene coordinates; not retroactively certified. Default browser/static new views use RefSeq.",
               "reference_url": config["reference_url"], "assembly_url": config["assembly_url"], "reference_download_date": config["reference_download_date"],
               "representative_region_bp": config["representative_region_bp"], "region_selection_rule": config["region_selection_rule"],
               "original_run_outputs_verified": checked, "training_performed": False, "raw_contact_scan_performed": False,
               "full_465_pngs_regenerated": False,
               "limitations": ["Reference annotation is a versioned snapshot, not experimentally verified strain-specific genes.",
                               "4296 exact closed-coordinate rows do not certify all supplied gene rows or the separate author structure convention.",
                               "Author labels are not classifier predictions; nested intervals are retained.",
                               "Original PNGs keep historical annotations and may have per-page y-scales; use the new browser for aligned RefSeq comparison.",
                               "Contact tracks are relative summaries; gene proximity is not a causal or new-structure claim."]}
    # No hard-coded matching count is used as an acceptance rule.
    summary["limitations"][1] = "Exact coordinate matches apply only to matched rows, not all supplied genes or the separate author structure convention."
    write_json(out/"report.json", summary)
    source_paths = [config_path, ROOT/config["reference_gff"], ROOT/config["assembly_report"], ROOT/config["author_workbook"],
                    ROOT/config["author_structures"], ROOT/config["source_gene_table"], run/"manifest.json", run/"task3/signals.npz", run/"task3/pages.csv",
                    Path(__file__), asset_dir/"atlas.js", asset_dir/"atlas.css"]
    write_json(out/"sources.json", {"files": [record(Path(p)) for p in source_paths]})
    write_json(out/"ARTIFACTS.json", {"files": [record(p, out) for p in sorted(out.rglob("*")) if p.is_file()]})
    write_json(out/"COMPLETED.json", {"status": "completed", "artifact_manifest_sha256": sha(out/"ARTIFACTS.json")})
    verify(out, require_local_inputs=True)
    write_json(report_dir/"review.json", {**summary, "delivery": str(out.relative_to(ROOT)), "sources": json.loads((out/"sources.json").read_text())})
    return summary


def verify(directory, *, require_local_inputs=False):
    out = Path(directory)
    done = json.loads((out/"COMPLETED.json").read_text())
    if done["artifact_manifest_sha256"] != sha(out/"ARTIFACTS.json"):
        raise ValueError("Artifact manifest changed")
    artifacts = json.loads((out/"ARTIFACTS.json").read_text())["files"]
    expected = {r["path"] for r in artifacts}
    actual = {str(p.relative_to(out)) for p in out.rglob("*") if p.is_file()} - {"ARTIFACTS.json", "COMPLETED.json"}
    if actual != expected:
        raise ValueError("Incomplete or extra atlas files")
    for r in artifacts:
        p = out/r["path"]
        if Path(r["path"]).is_absolute() or ".." in Path(r["path"]).parts or p.is_symlink() or sha(p) != r["sha256"]:
            raise ValueError("Atlas payload changed")
    report = json.loads((out/"report.json").read_text())
    pages = rows(out/"pages.csv")
    check_pages(pages, out, report["genome_length_bp"], 10_000)
    for r in pages:
        if not (out/r["html"]).is_file():
            raise ValueError("Missing HTML page")
    parsed = read_reference(out/"reference_genomic.gff.gz", "NC_000913.3", report["genome_length_bp"])
    saved = rows(out/"reference_genes.csv")
    if [{k: str(v) for k, v in r.items()} for r in parsed] != saved:
        raise ValueError("Reference conversion mismatch")
    gene_audit = rows(out/"supplied_gene_audit.csv")
    restored_source = [{"gene": r["gene"], "start": r["source_start"], "end": r["source_end"],
                        "strand": r["source_strand"]} for r in gene_audit]
    recomputed = [{k: str(v) for k, v in r.items()} for r in audit_gene_rows(restored_source, parsed)]
    if recomputed != gene_audit:
        raise ValueError("Gene coordinate comparison does not recompute")
    if dict(Counter(r["status"] for r in gene_audit)) != report["supplied_gene_coordinate_comparison"]:
        raise ValueError("Annotation audit count differs")
    if report["representative_region_bp"] != [1_050_000, 1_100_000]:
        raise ValueError("Fixed representative region changed")
    for r in json.loads((out/"sources.json").read_text())["files"]:
        p = ROOT/r["path"]
        if p.is_file():
            if sha(p) != r["sha256"]:
                raise ValueError(f"Source changed: {p}")
        elif require_local_inputs:
            raise FileNotFoundError(p)
    return {"status": "verified", "pages": len(pages), "actual_pngs_checked": len(pages),
            "reference_gene_features": len(parsed), "artifact_files": len(artifacts), "training_performed": False}


def page_html(i, page):
    page = {**page, "start": int(page["start"]), "end": int(page["end"])}
    return f'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Genome tracks {page['start']}–{page['end']}</title><link rel="stylesheet" href="../atlas.css"><body data-page="{i}" data-base="../"><main><a href="../index.html">← 三级图集总览</a><h1>{page['start']:,}–{page['end']:,} bp</h1><div id="browser"></div><noscript><p>交互RefSeq注释需要JavaScript。下图为保留的作者注释版本，坐标约定仍暂定。</p><img src="../{page['file']}" alt="原始轨道"></noscript></main><script src="../atlas_data.js"></script><script src="../atlas.js"></script></body></html>'''


def index_html(pages):
    links = "\n".join(f'<a href="{r["html"]}" data-page-link="{i}">{i+1:03d} · {int(r["start"]):,}–{int(r["end"]):,}</a>' for i, r in enumerate(pages))
    return f'''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>MicroCTure | Genome track atlas</title><link rel="stylesheet" href="atlas.css"><body data-base=""><main><header><p class="eyebrow">MICROCTURE / TASK 3 / DIGITAL SUPPLEMENT</p><h1>全基因组多条件接触轨道</h1><p class="lead">4.64 Mb · 3 conditions · 6 samples · 465 complete track pages</p><nav><a href="#whole">01 全基因组</a><a href="#browse">02 465页浏览</a><a href="#detail">03 固定代表区域</a><a href="#provenance">注释与来源</a></nav></header>
<section id="whole"><p class="eyebrow">LEVEL 1 / WHOLE-GENOME OVERVIEW</p><h2>整个4,641,652 bp染色体，同一坐标与条件量纲</h2><p>10 kb显示聚合，保留100 bp数值；细线为重复、粗线为条件均值。作者结构与RefSeq正负链基因独立成轨，灰带标出固定代表区域。</p><a href="whole_genome.png"><img src="whole_genome.png" alt="六轨全基因组总览" loading="eager"></a></section>
<section id="browse"><p class="eyebrow">LEVEL 2 / COMPLETE DIGITAL BROWSER</p><h2>465页完整浏览，不以精选图替代全基因组</h2><div id="browser"></div><details><summary>全部465页坐标索引</summary><div class="page-grid">{links}</div></details></section>
<section id="detail"><p class="eyebrow">LEVEL 3 / PREDEFINED REPRESENTATIVE REGION</p><h2>固定1.05–1.10 Mb：从总览进入注释细节</h2><p>The representative region was predefined independently of classification performance and condition-specific effect size.</p><p>该区域沿用既有固定选择规则，没有按本轮差异或分类表现重新选位。基因箭头使用独立核验的RefSeq坐标；作者区间保留原标注及嵌套。</p><a href="representative_region.png"><img src="representative_region.png" alt="预先固定的1.05–1.10 Mb局部多轨图" loading="lazy"></a></section>
<section id="provenance"><p class="eyebrow">ANNOTATION PROVENANCE</p><h2>区分参考基因、作者区间和原始绘图注释</h2><p>新总览、固定区域及默认交互视图：NCBI RefSeq GCF_000005845.2 / NC_000913.3，4,651条gene/pseudogene记录；GFF由1-based闭区间明确转为0-based半开区间。提供完整固定快照、逐行比较及SHA-256来源。</p><p>作者补充表4/5/6的344结构已逐条核对原工作簿；原作者坐标的端点约定仍暂定，不把来源一致性称作生物学真值认证。补充表1的4,516行与参考存在版本、命名和区间差异，逐行保留，不自动改标或删除。</p><p class="notice">保留的465张原始PNG仍使用作者补充表基因原坐标；不是重新绘制的RefSeq版本。浏览器默认显示经过坐标转换的独立RefSeq轨。原图按页可下载，二者清楚区分。</p><p>轨道是归一化接触摘要，不是分类预测、显著性检验或基因因果作用证据。两个重复和相邻窗口不增加独立生物学样本数。</p><nav><a href="reference_genes.csv">参考基因表</a><a href="supplied_gene_audit.csv">4,516行比较</a><a href="author_structure_audit.csv">344结构来源核验</a><a href="reference_genomic.gff.gz">NCBI GFF快照</a><a href="assembly_report.txt">组装说明</a><a href="pages.csv">全页索引</a><a href="report.json">核验结果</a><a href="sources.json">来源签名</a></nav><p><a href="https://www.ncbi.nlm.nih.gov/datasets/genome/GCF_000005845.2/">NCBI reference assembly</a> · <a href="https://doi.org/10.1038/s41586-025-09396-y">作者原研究</a></p></section></main><script src="atlas_data.js"></script><script src="atlas.js"></script></body></html>'''


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=("build", "verify"))
    parser.add_argument("--config", type=Path, default=CONFIG)
    parser.add_argument("--directory", type=Path, default=ROOT/"submission/task3/atlas/v1")
    parser.add_argument("--require-local-inputs", action="store_true")
    args = parser.parse_args(argv)
    result = build(args.config) if args.action == "build" else verify(args.directory, require_local_inputs=args.require_local_inputs)
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
